"""Word (.docx) report — builds without a live LLM (structured fallback)."""
import io

from docx import Document

from app.pcap.capture import Capture
from app.pcap.docx_report import build_docx


def _capture(sample_pcap_metadata, sample_pcap_agents) -> Capture:
    return Capture(
        id="cap123",
        metadata=sample_pcap_metadata,
        agents=sample_pcap_agents,
        filename="capture.pcap",
    )


def _headings(doc: Document) -> list[str]:
    return [p.text for p in doc.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_docx_builds_without_llm(pcap_db, sample_pcap_metadata, sample_pcap_agents, monkeypatch):
    import app.pcap.docx_report as dr
    monkeypatch.setattr(dr, "generate_report", lambda c: (_ for _ in ()).throw(RuntimeError("no llm")))

    buf = build_docx(_capture(sample_pcap_metadata, sample_pcap_agents))
    assert isinstance(buf, io.BytesIO)
    doc = Document(buf)
    headings = _headings(doc)
    assert "PCAP Security Incident Report" in headings
    assert "Capture Summary" in headings
    assert "Agent Findings" in headings
    assert any("Threat Detection Agent" in h for h in headings)


def test_docx_with_edge_context_and_chat(pcap_db, sample_pcap_metadata, sample_pcap_agents, monkeypatch):
    import app.pcap.docx_report as dr
    monkeypatch.setattr(dr, "generate_report", lambda c: "## Summary\nAll clear.")

    edge_event = {
        "scenario_id": "c2_beacon",
        "pcap_source": "bundled",
        "staging_backend": "local",
        "staged_uri": "/tmp/x.pcap",
        "edge_actions": [{"action": "block_ip", "target": "45.133.1.77", "detail": "blackhole"}],
        "site": {"name": "Denver Headend", "location": "HE-TEST-01", "model_name": "edge-v2"},
    }
    transcript = [
        {"role": "user", "content": "Is this beaconing?"},
        {"role": "assistant", "content": "Yes, periodic outbound."},
    ]
    buf = build_docx(
        _capture(sample_pcap_metadata, sample_pcap_agents),
        edge_event=edge_event,
        chat_transcript=transcript,
    )
    doc = Document(buf)
    headings = _headings(doc)
    assert "Edge Data Center Context" in headings
    assert "Operator Chat Transcript" in headings
    text = _all_text(doc)
    assert "Denver Headend" in text
    assert "Is this beaconing?" in text


def test_docx_endpoint(pcap_client, sample_pcap_metadata, sample_pcap_agents, monkeypatch):
    # Persist a capture through the store so the endpoint can find it.
    from app.pcap import analysis
    from app.database import SessionLocal

    cap = _capture(sample_pcap_metadata, sample_pcap_agents)
    cap.id = "apicap"
    with SessionLocal() as db:
        analysis.persist_capture(db, cap)

    import app.pcap.docx_report as dr
    monkeypatch.setattr(dr, "generate_report", lambda c: "## Summary\nnothing notable")

    resp = pcap_client.post("/api/v1/pcap/report/docx", json={"capture_id": "apicap", "include_chat": False})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    Document(io.BytesIO(resp.content))


def test_docx_missing_capture_404(pcap_client):
    resp = pcap_client.post("/api/v1/pcap/report/docx", json={"capture_id": "nope"})
    assert resp.status_code == 404
