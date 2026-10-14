"""Word (.docx) incident report builder.

Assembles a formatted report from a stored analysis session (and optional edge
context + chat transcript). The executive summary reuses the existing LLM
markdown report when available; if no LLM is reachable it degrades to a
structured-only summary so docx export always works.
"""
from __future__ import annotations

import io
import logging
import re
from datetime import datetime, timezone

from docx import Document
from docx.shared import Pt, RGBColor

from .capture import Capture as Session
from .report import generate_report
from .schemas import AgentResponse, PcapMetadata

log = logging.getLogger(__name__)

_SEV_ORDER = ["low", "medium", "high", "critical"]
_SEV_RGB = {
    "low": RGBColor(0x2E, 0x7D, 0x32),
    "medium": RGBColor(0xB8, 0x8A, 0x00),
    "high": RGBColor(0xD8, 0x4B, 0x00),
    "critical": RGBColor(0xC6, 0x28, 0x28),
}


def _highest_severity(agents: list[AgentResponse]) -> str:
    highest = "low"
    for a in agents:
        if _SEV_ORDER.index(a.severity) > _SEV_ORDER.index(highest):
            highest = a.severity
    return highest


def _add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
        for para in cells[0].paragraphs:
            for run in para.runs:
                run.bold = True


def _markdown_to_docx(doc: Document, markdown: str) -> None:
    """Minimal markdown → docx walker: headings, bullets, bold runs."""
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif re.match(r"^\s*[-*]\s+", line):
            text = re.sub(r"^\s*[-*]\s+", "", line)
            _add_rich_paragraph(doc, text, style="List Bullet")
        elif re.match(r"^\s*\d+\.\s+", line):
            text = re.sub(r"^\s*\d+\.\s+", "", line)
            _add_rich_paragraph(doc, text, style="List Number")
        else:
            _add_rich_paragraph(doc, line)


def _add_rich_paragraph(doc: Document, text: str, style: str | None = None) -> None:
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    # Split on **bold** spans.
    for i, chunk in enumerate(re.split(r"\*\*(.+?)\*\*", text)):
        if not chunk:
            continue
        run = para.add_run(chunk)
        if i % 2 == 1:  # odd chunks were inside ** **
            run.bold = True


def build_docx(
    session: Session,
    *,
    edge_event: dict | None = None,
    chat_transcript: list[dict] | None = None,
) -> io.BytesIO:
    doc = Document()
    meta: PcapMetadata = session.metadata
    agents = session.agents
    highest = _highest_severity(agents)

    # --- Title page -------------------------------------------------------
    title = doc.add_heading("PCAP Security Incident Report", level=0)
    subtitle = doc.add_paragraph()
    run = subtitle.add_run(session.filename or "packet capture")
    run.italic = True
    run.font.size = Pt(12)

    sev_para = doc.add_paragraph()
    sev_para.add_run("Highest severity: ").bold = True
    sev_run = sev_para.add_run(highest.upper())
    sev_run.bold = True
    sev_run.font.color.rgb = _SEV_RGB[highest]

    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    doc.add_paragraph(f"Generated: {generated}")

    if edge_event:
        site = edge_event.get("site") or {}
        doc.add_heading("Edge Data Center Context", level=1)
        _add_kv_table(
            doc,
            [
                ("Edge site", site.get("name", "—")),
                ("Location", site.get("location", "—")),
                ("Local model", site.get("model_name", "—")),
                ("Scenario", edge_event.get("scenario_id", "—")),
                ("Capture source", edge_event.get("pcap_source", "—")),
                ("Staging backend", edge_event.get("staging_backend", "—")),
                ("Staged URI", edge_event.get("staged_uri") or "—"),
            ],
        )
        actions = edge_event.get("edge_actions") or []
        if actions:
            doc.add_heading("Actions taken at the edge", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "Action", "Target", "Detail"
            for a in actions:
                cells = table.add_row().cells
                cells[0].text = str(a.get("action", ""))
                cells[1].text = str(a.get("target", ""))
                cells[2].text = str(a.get("detail", ""))

    # --- Capture summary --------------------------------------------------
    doc.add_heading("Capture Summary", level=1)
    _add_kv_table(
        doc,
        [
            ("Total packets", f"{meta.total_packets:,}"),
            ("Time window", f"{meta.time_range_start} → {meta.time_range_end}"),
            ("Unique source IPs", str(meta.unique_src_ips)),
            ("Unique destination IPs", str(meta.unique_dst_ips)),
            ("Protocols", ", ".join(f"{k} ({v})" for k, v in meta.protocols.items()) or "—"),
            ("Transport flows", str(len(getattr(meta, "flows", []) or []))),
            ("DNS queries", str(len(getattr(meta, "dns_queries", []) or []))),
            ("TLS handshakes", str(len(getattr(meta, "tls_handshakes", []) or []))),
            ("External IPs (geo-located)", str(len(meta.geo_locations or {}))),
            ("Payload findings", str(len(meta.payload_findings))),
            ("HTTP requests", str(len(meta.http_requests))),
        ],
    )

    # --- Executive summary (LLM, with structured fallback) ----------------
    doc.add_heading("Executive Summary", level=1)
    try:
        markdown = generate_report(session)
        _markdown_to_docx(doc, markdown)
    except Exception as exc:  # noqa: BLE001
        log.warning("docx: LLM report unavailable (%s); using structured fallback", exc)
        doc.add_paragraph(
            "An AI-authored narrative could not be generated (no LLM reachable). The "
            "structured agent findings below are included in full."
        )

    # --- Per-agent sections ----------------------------------------------
    doc.add_heading("Agent Findings", level=1)
    for agent in agents:
        heading = doc.add_heading(level=2)
        run = heading.add_run(f"{agent.agent_name} — {agent.severity.upper()}")
        run.font.color.rgb = _SEV_RGB.get(agent.severity, _SEV_RGB["low"])

        if agent.analysis:
            doc.add_paragraph(agent.analysis)

        if agent.findings:
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "Type", "Description", "Risk"
            for f in agent.findings:
                cells = table.add_row().cells
                cells[0].text = f.type
                cells[1].text = f.description
                cells[2].text = f.risk or "—"

        if agent.recommendations:
            doc.add_heading("Recommendations", level=3)
            for rec in agent.recommendations:
                doc.add_paragraph(rec, style="List Bullet")

    # --- Threat-intel appendix -------------------------------------------
    blocklisted = [r for r in (meta.ip_reputations or {}).values() if r.in_blocklist]
    if blocklisted or meta.cve_candidates:
        doc.add_heading("Threat Intelligence Appendix", level=1)
        if blocklisted:
            doc.add_heading("Blocklisted IPs", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "IP", "Sources", "Categories"
            for r in blocklisted:
                cells = table.add_row().cells
                cells[0].text = r.ip
                cells[1].text = ", ".join(r.sources) or "—"
                cells[2].text = ", ".join(r.categories) or "—"
        if meta.cve_candidates:
            doc.add_heading("CVE Candidates", level=2)
            table = doc.add_table(rows=1, cols=3)
            table.style = "Light Grid Accent 1"
            hdr = table.rows[0].cells
            hdr[0].text, hdr[1].text, hdr[2].text = "CVE", "Severity", "Summary"
            for cve in meta.cve_candidates:
                cells = table.add_row().cells
                cells[0].text = cve.id
                cells[1].text = cve.severity or "—"
                cells[2].text = cve.summary

    # --- Chat appendix ----------------------------------------------------
    if chat_transcript:
        doc.add_heading("Operator Chat Transcript", level=1)
        for msg in chat_transcript:
            role = str(msg.get("role", "")).capitalize()
            content = str(msg.get("content", ""))
            if not content.strip():
                continue
            para = doc.add_paragraph()
            para.add_run(f"{role}: ").bold = True
            para.add_run(content)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
